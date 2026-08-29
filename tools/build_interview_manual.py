from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
PALE_BLUE = "E8EEF5"
LIGHT_BLUE = "F4F7FA"
GRAY = "666666"
LIGHT_GRAY = "F3F4F6"
WHITE = "FFFFFF"
CONTENT_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    # Left alignment plus an explicit indent keeps Word from dropping tblInd
    # when it normalizes the package. The indent matches the cells' start
    # margin, so the visible border aligns with surrounding body text.
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = borders.find(qn(f"w:{side}"))
        if edge is None:
            edge = OxmlElement(f"w:{side}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "4")
        edge.set(qn("w:color"), "C9D2DC")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_field(run, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, end))


def add_hyperlink(paragraph, text: str, url: str):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend((color, underline))
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`|https?://[^\s]+)")


def add_inline(paragraph, text: str, base_size: float = 11.0, color: str | None = None) -> None:
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            run.font.size = Pt(base_size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.font.size = Pt(base_size)
            if color:
                run.font.color.rgb = RGBColor.from_string(color)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(max(base_size - 0.5, 8.5))
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            r_pr = run._r.get_or_add_rPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), LIGHT_GRAY)
            r_pr.append(shd)
        else:
            clean = token.rstrip(".,;，。；）)")
            suffix = token[len(clean):]
            add_hyperlink(paragraph, clean, clean)
            if suffix:
                paragraph.add_run(suffix)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.size = Pt(base_size)
        if color:
            run.font.color.rgb = RGBColor.from_string(color)


def set_east_asia_font(run, name="Microsoft YaHei") -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def style_all_runs(paragraph, east_asia="Microsoft YaHei") -> None:
    for run in paragraph.runs:
        set_east_asia_font(run, east_asia)


def add_quote(document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.style = document.styles["Quote"]
    paragraph.paragraph_format.left_indent = Inches(0.22)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.22
    add_inline(paragraph, "\n".join(lines), 10.5, DARK_BLUE)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_BLUE)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), BLUE)
    borders.append(left)
    p_pr.append(borders)
    style_all_runs(paragraph)


def add_code(document, lines: list[str]) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.1)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.8)
    run.font.color.rgb = RGBColor.from_string("273444")
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    for side in ("top", "left", "bottom", "right"):
        edge = OxmlElement(f"w:{side}")
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "3")
        edge.set(qn("w:color"), "D8DEE5")
        borders.append(edge)
    p_pr.append(borders)


def widths_for_table(col_count: int) -> list[int]:
    if col_count == 2:
        return [2600, 6760]
    if col_count == 4:
        return [1350, 2350, 2350, 3310]
    if col_count == 5:
        return [700, 1500, 2350, 2200, 2610]
    base = CONTENT_DXA // col_count
    widths = [base] * col_count
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_markdown_table(document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=col_count)
    table.style = "Table Grid"
    widths = widths_for_table(col_count)
    set_table_geometry(table, widths)
    set_repeat_table_header(table.rows[0])
    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.12
            add_inline(paragraph, value.strip(), 9.0 if col_count >= 4 else 9.5)
            style_all_runs(paragraph)
            if r_idx == 0:
                set_cell_shading(cell, PALE_BLUE)
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, "F8FAFC")
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_toc(document) -> None:
    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    title_run = title.add_run("目录")
    title_run.bold = True
    title_run.font.size = Pt(16)
    title_run.font.color.rgb = RGBColor.from_string(BLUE)
    set_east_asia_font(title_run)
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    add_field(run, 'TOC \\o "1-1" \\h \\z \\u')


def configure_styles(document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    styles["Heading 1"].paragraph_format.page_break_before = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    quote = styles["Quote"]
    quote.font.name = "Calibri"
    quote._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    quote.font.size = Pt(10.5)
    quote.font.italic = False


def configure_section(document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.42)


def add_header_footer(section) -> None:
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""
    header = section.header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("星演票务项目面试学习手册  ·  袁祥凯")
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(8.5)
    run.font.color.rgb = RGBColor.from_string(GRAY)
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "C9D2DC")
    borders.append(bottom)
    p_pr.append(borders)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("—  ")
    r1.font.color.rgb = RGBColor.from_string("A0A0A0")
    r2 = p.add_run()
    add_field(r2, "PAGE")
    r3 = p.add_run("  —")
    for run in (r1, r2, r3):
        run.font.name = "Calibri"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_cover(document) -> None:
    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(48)
    p.paragraph_format.space_after = Pt(14)
    run = p.add_run("JAVA BACKEND · PROJECT INTERVIEW GUIDE")
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run("星演票务项目")
    run.font.name = "Calibri Light"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(30)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run("简历亮点 · 技术学习 · 真实面经逐字稿")
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(17)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(22)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:color"), BLUE)
    borders.append(bottom)
    p_pr.append(borders)

    for label, value in (
        ("候选人", "袁祥凯｜电子科技大学｜2027 届"),
        ("代码基线", "stellaris-platform · v5/reference 主链路"),
        ("文档目标", "把每一条简历亮点讲到正常链路、失败窗口与工程边界"),
        ("复核日期", "2026 年 8 月 28 日"),
    ):
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r1 = p.add_run(f"{label}  ")
        r1.bold = True
        r1.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        r2 = p.add_run(value)
        r2.font.color.rgb = RGBColor.from_string("3E4A59")
        for run in p.runs:
            set_east_asia_font(run)
            run.font.size = Pt(11)

    p = document.add_paragraph()
    p.paragraph_format.space_before = Pt(26)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("核心口径")
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    p = document.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.right_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    run = p.add_run("至少一次投递 + 业务幂等 + 数据库 CAS + 最终一致；不宣称 Exactly Once、零丢失或未经验证的 QPS。")
    set_east_asia_font(run)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), PALE_BLUE)
    p_pr.append(shd)

    document.add_page_break()


def parse_table(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def build(source: Path, output: Path) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_section(document)
    configure_styles(document)
    add_header_footer(document.sections[0])
    settings = document.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    add_cover(document)

    idx = 0
    in_code = False
    code_lines: list[str] = []
    quote_lines: list[str] = []
    toc_added = False

    while idx < len(lines):
        raw = lines[idx]
        line = raw.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if stripped.startswith(">"):
            quote_lines.append(stripped[1:].lstrip())
            idx += 1
            while idx < len(lines) and lines[idx].strip().startswith(">"):
                quote_lines.append(lines[idx].strip()[1:].lstrip())
                idx += 1
            add_quote(document, quote_lines)
            quote_lines = []
            continue

        if stripped.startswith("|"):
            rows, idx = parse_table(lines, idx)
            add_markdown_table(document, rows)
            continue

        if not stripped or stripped == "---":
            idx += 1
            continue

        if stripped.startswith("# "):
            idx += 1
            continue

        if stripped.startswith("## "):
            text = stripped[3:].strip()
            if text.startswith("2.") and not toc_added:
                document.add_page_break()
                add_toc(document)
                toc_added = True
            p = document.add_paragraph()
            p.style = document.styles["Heading 1"]
            if text.startswith("0.") or (text.startswith("2.") and toc_added):
                p.paragraph_format.page_break_before = False
            add_inline(p, text, 16, BLUE)
            style_all_runs(p)
            idx += 1
            continue

        if stripped.startswith("### "):
            p = document.add_paragraph()
            p.style = document.styles["Heading 2"]
            add_inline(p, stripped[4:].strip(), 13, BLUE)
            style_all_runs(p)
            idx += 1
            continue

        if stripped.startswith("#### "):
            p = document.add_paragraph()
            p.style = document.styles["Heading 3"]
            add_inline(p, stripped[5:].strip(), 12, DARK_BLUE)
            style_all_runs(p)
            idx += 1
            continue

        bullet = re.match(r"^[-*]\s+(.+)$", stripped)
        number = re.match(r"^\d+\.\s+(.+)$", stripped)
        if bullet or number:
            p = document.add_paragraph()
            if bullet:
                p.style = document.styles["List Bullet"]
                item_text = bullet.group(1)
            else:
                p.paragraph_format.left_indent = Inches(0.375)
                p.paragraph_format.first_line_indent = Inches(-0.188)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.25
                item_text = f"{stripped.split('.', 1)[0]}. {number.group(1)}"
            add_inline(p, item_text, 11)
            style_all_runs(p)
            idx += 1
            continue

        p = document.add_paragraph()
        p.paragraph_format.widow_control = True
        add_inline(p, stripped.rstrip("  "), 11)
        style_all_runs(p)
        idx += 1

    if in_code and code_lines:
        add_code(document, code_lines)

    core = document.core_properties
    core.title = "星演票务项目：简历亮点、技术学习与真实面经逐字稿"
    core.subject = "Java 后端项目面试学习手册"
    core.author = "袁祥凯"
    core.keywords = "Java, Redis, Lua, Kafka, MySQL, 高并发, 面试"
    core.comments = "基于 stellaris-platform 当前代码复核生成"

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source, args.output)


if __name__ == "__main__":
    main()
